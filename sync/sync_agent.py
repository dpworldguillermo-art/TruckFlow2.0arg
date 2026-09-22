import json, time, sys
from pathlib import Path
from datetime import datetime
import requests
from openpyxl import load_workbook

HEADER_KEY='Transaction Type'

def read_config(path):
    with open(path,'r',encoding='utf-8') as f: return json.load(f)

def newest_excel(folder):
    files=[p for p in Path(folder).glob('*.xlsx') if not p.name.startswith('~$')]
    return max(files,key=lambda p:p.stat().st_mtime) if files else None

def cell_value(v):
    if isinstance(v,datetime): return v.strftime('%Y-%m-%d %H:%M:%S')
    return v

def parse_excel(path):
    wb=load_workbook(path,read_only=True,data_only=True)
    ws=wb[wb.sheetnames[0]]
    header_row=None; headers=None
    for i,row in enumerate(ws.iter_rows(values_only=True),start=1):
        vals=list(row)
        if HEADER_KEY in vals:
            header_row=i; headers=[str(v).strip() if v is not None else '' for v in vals]; break
    if not header_row: raise RuntimeError('No se encontró la fila de encabezados con Transaction Type')
    rows=[]
    for row in ws.iter_rows(min_row=header_row+1,values_only=True):
        if not any(v is not None and str(v).strip() for v in row): continue
        rec={headers[i]:cell_value(row[i]) if i<len(row) else None for i in range(len(headers)) if headers[i]}
        rows.append(rec)
    return {'sourceFile':path.name,'generatedAt':datetime.now().astimezone().isoformat(),'count':len(rows),'rows':rows}

def publish(payload,cfg):
    r=requests.post(cfg['endpoint'],json=payload,headers={'Authorization':f"Bearer {cfg['token']}"},timeout=60)
    r.raise_for_status()

def main():
    cfg=read_config(sys.argv[1] if len(sys.argv)>1 else 'config.json')
    last=None; interval=int(cfg.get('interval_seconds',10))
    print('Gambetta Sync iniciado. Carpeta:',cfg['folder'])
    while True:
        try:
            f=newest_excel(cfg['folder'])
            if f:
                stamp=(str(f.resolve()),f.stat().st_mtime_ns,f.stat().st_size)
                if stamp!=last:
                    payload=parse_excel(f); publish(payload,cfg); last=stamp
                    print(datetime.now().strftime('%H:%M:%S'),'Publicado',f.name,'-',payload['count'],'registros')
        except Exception as e:
            print(datetime.now().strftime('%H:%M:%S'),'ERROR:',e)
        time.sleep(interval)

if __name__=='__main__': main()
